from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

from bling_app_zero.core.audit import add_audit_event

RESPONSIBLE_FILE = 'bling_app_zero/core/excel_running_guard.py'
MAX_HEADER_SCAN_COLUMNS = 512
MAX_HEADER_SCAN_ROWS = 80
MAX_HEADER_SCAN_SHEETS = 20
MAX_DATA_ROWS = 100_000
BLANK_ROW_STOP = 500
PATCH_ATTR = '_mapeiaai_excel_running_guard_v1'


def _bounded_find_best_excel_header(workbook: Any):
    from bling_app_zero.core import files as files_module

    candidates = []
    sheets = list(getattr(workbook, 'worksheets', []) or [])[:MAX_HEADER_SCAN_SHEETS]
    for sheet_index, sheet in enumerate(sheets):
        max_row = min(int(getattr(sheet, 'max_row', 0) or 0), MAX_HEADER_SCAN_ROWS)
        max_col = min(int(getattr(sheet, 'max_column', 0) or 0), MAX_HEADER_SCAN_COLUMNS)
        if max_row <= 0 or max_col <= 0:
            continue
        for row_index in range(1, max_row + 1):
            candidate = files_module._score_header_candidate(sheet, sheet_index, row_index, max_col)
            if candidate is not None:
                candidates.append(candidate)
    if not candidates:
        return None
    candidates.sort(
        key=lambda item: (item.score, item.data_rows_below, len(item.columns), -item.header_row),
        reverse=True,
    )
    return candidates[0]


def _bounded_frame_from_excel_candidate(workbook: Any, candidate: Any) -> pd.DataFrame:
    from bling_app_zero.core import files as files_module

    sheet = workbook.worksheets[candidate.sheet_index]
    positions = [int(position) for position in candidate.positions if 0 < int(position) <= MAX_HEADER_SCAN_COLUMNS]
    if not positions:
        return pd.DataFrame(columns=list(candidate.columns))

    first_row = int(candidate.header_row) + 1
    declared_last_row = int(getattr(sheet, 'max_row', first_row) or first_row)
    last_row = min(declared_last_row, first_row + MAX_DATA_ROWS - 1)
    max_needed_column = max(positions)
    rows: list[list[str]] = []
    consecutive_blank = 0

    iterator = sheet.iter_rows(
        min_row=first_row,
        max_row=last_row,
        min_col=1,
        max_col=max_needed_column,
        values_only=True,
    )
    for raw_row in iterator:
        row = [files_module._clean(raw_row[position - 1] if position - 1 < len(raw_row) else '') for position in positions]
        if any(files_module._clean(value) for value in row):
            rows.append(row)
            consecutive_blank = 0
            continue
        consecutive_blank += 1
        if consecutive_blank >= BLANK_ROW_STOP:
            break

    return files_module._clean_columns(pd.DataFrame(rows, columns=list(candidate.columns)).fillna(''))


def _bounded_excel_header_columns(template_name: str, template_bytes: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except Exception:
        return []

    workbook = None
    try:
        keep_vba = Path(template_name).suffix.lower() == '.xlsm'
        workbook = load_workbook(BytesIO(template_bytes), read_only=True, data_only=False, keep_vba=keep_vba)
        best: list[str] = []
        sheets = list(getattr(workbook, 'worksheets', []) or [])[:MAX_HEADER_SCAN_SHEETS]
        for sheet in sheets:
            max_rows = min(int(getattr(sheet, 'max_row', 1) or 1), 30)
            max_cols = min(int(getattr(sheet, 'max_column', 0) or 0), MAX_HEADER_SCAN_COLUMNS)
            if max_cols <= 0:
                continue
            for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
                columns = [' '.join(str(value or '').replace('\ufeff', '').replace('\x00', '').replace('\xa0', ' ').split()).strip() for value in row]
                columns = [value for value in columns if value]
                if len(columns) > len(best):
                    best = columns
        return best
    except Exception:
        return []
    finally:
        try:
            if workbook is not None:
                workbook.close()
        except Exception:
            pass


def install_excel_running_guard() -> bool:
    from bling_app_zero.core import files as files_module
    from bling_app_zero.core import final_template_exporter

    if getattr(files_module, PATCH_ATTR, False):
        return False

    files_module._find_best_excel_header = _bounded_find_best_excel_header
    files_module._frame_from_excel_candidate = _bounded_frame_from_excel_candidate
    final_template_exporter._excel_header_columns = _bounded_excel_header_columns
    setattr(files_module, PATCH_ATTR, True)

    add_audit_event(
        'excel_running_guard_installed',
        area='MODELO',
        status='OK',
        details={
            'max_header_columns': MAX_HEADER_SCAN_COLUMNS,
            'max_header_rows': MAX_HEADER_SCAN_ROWS,
            'max_data_rows': MAX_DATA_ROWS,
            'blank_row_stop': BLANK_ROW_STOP,
            'prevents_formatted_empty_rows_hang': True,
            'responsible_file': RESPONSIBLE_FILE,
        },
    )
    return True


__all__ = ['install_excel_running_guard']
